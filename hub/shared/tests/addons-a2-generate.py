"""Deterministic synthetic importer fixtures; no research data or estimates."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED, ZIP_STORED
from datetime import datetime
import re
import openpyxl

ROOT = Path(__file__).with_name('addons-a2-fixtures')
ROOT.mkdir(exist_ok=True)

def parts(wb):
    wb.properties.created = datetime(2000, 1, 1)
    wb.properties.modified = datetime(2000, 1, 1)
    buf = BytesIO()
    wb.save(buf)
    with ZipFile(buf) as z:
        return {n: re.sub(rb'<dcterms:(created|modified)[^>]*>.*?</dcterms:\1>', lambda m: b'<dcterms:'+m[1]+b' xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:type="dcterms:W3CDTF">2000-01-01T00:00:00Z</dcterms:'+m[1]+b'>', z.read(n)) if n == 'docProps/core.xml' else z.read(n) for n in z.namelist()}

def write(name, data, method=ZIP_DEFLATED):
    with ZipFile(ROOT / name, 'w') as z:
        for path, value in sorted(data.items()):
            info = ZipInfo(path, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = method
            z.writestr(info, value)

def altered(data, path, old, new):
    result = data.copy()
    assert old in result[path]
    result[path] = result[path].replace(old, new)
    return result

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Precision'
ws.append(['study', 'effect', 'se'])
raw = ['0.12345678901234567', '1.0000000000000002', '5e-324', '1.7976931348623157e308', '-2.2250738585072014e-308', '9007199254740991']
for i, value in enumerate(raw, 2):
    ws.append([f'Fixture {i-1}', 1, 0.1])
    ws.cell(i, 2).number_format = '0.00'
base = parts(wb)
path = 'xl/worksheets/sheet1.xml'
xml = base[path].decode()
for i, value in enumerate(raw, 2):
    xml = re.sub(fr'(<c r="B{i}"[^>]*><v>).*?(</v>)', lambda m: m[1]+value+m[2], xml)
base[path] = xml.encode()
write('precision.xlsx', base)
write('stored.xlsx', base, ZIP_STORED)
(ROOT / 'precision.csv').write_text('study,effect,se\n'+'\n'.join(f'Fixture {i},{v},0.1' for i,v in enumerate(raw,1)), encoding='utf-8')

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Warnings'
ws.append(['study', 'value', 'extra'])
ws.append(['Comma, quote " and\nnewline', 0.125, '=1+1']); ws['B2'].number_format='0.0%'
ws.append(['Date fixture', datetime(2020,1,2), '=2+2'])
ws.append(['Merged', 3]); ws.merge_cells('A4:A5')
ws['C6'] = True
wb.create_sheet('Other').append(['alternative', 42])
warning = parts(wb)
warning[path] = warning[path].replace(b'<f>2+2</f><v></v>', b'<f>2+2</f><v>4</v>')
write('warnings.xlsx', warning)
# Hand-authored shared strings part supplements the openpyxl fixture.
shared = base.copy()
shared[path] = shared[path].replace(b'<c r="A2" t="inlineStr"><is><t>Fixture 1</t></is></c>', b'<c r="A2" t="s"><v>0</v></c>')
shared['xl/sharedStrings.xml'] = b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><r><t>Rich </t></r><r><t>text</t></r><rPh sb="0" eb="1"><t>ignored</t></rPh></si></sst>'
shared['xl/_rels/workbook.xml.rels'] = shared['xl/_rels/workbook.xml.rels'].replace(b'</Relationships>', b'<Relationship Id="strings" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
write('shared.xlsx', shared)
mutations = {
 'doctype': (path, b'<worksheet ', b'<!DOCTYPE worksheet [<!ENTITY x "boom">]><worksheet '),
 'malformed': (path, b'</worksheet>', b'</broken>'),
 'numeric': (path, raw[0].encode(), b'NaN'),
 'error-cell': (path, b'<c r="B2" s="1" t="n">', b'<c r="B2" s="1" t="e">'),
 'sparse-limit': (path, b'r="B2"', b'r="XFD1048576"'),
 'duplicate-cell': (path, b'r="B2"', b'r="A2"'),
 'style-index': (path, b's="1"', b's="999"'),
 'missing-part': ('xl/_rels/workbook.xml.rels', b'/xl/worksheets/sheet1.xml', b'/xl/worksheets/missing.xml'),
 'external': ('xl/_rels/workbook.xml.rels', b'Target="/xl/worksheets/sheet1.xml"', b'TargetMode="External" Target="https://example.invalid/sheet.xml"'),
}
for name, (part, old, new) in mutations.items():
    write(name+'.xlsx', altered(base, part, old, new))
write('bad-shared.xlsx', altered(shared, path, b'<c r="A2" t="s"><v>0</v>', b'<c r="A2" t="s"><v>9</v>'))
write('traversal.xlsx', {**base, '../bad.xml': b'<bad/>'})
# Corrupt payload while preserving ZIP metadata: CRC must catch this.
corrupt = bytearray((ROOT/'stored.xlsx').read_bytes())
pos = corrupt.index(b'Fixture 1'); corrupt[pos] = ord('X')
(ROOT/'crc.xlsx').write_bytes(corrupt)
(ROOT/'truncated.xlsx').write_bytes((ROOT/'precision.xlsx').read_bytes()[:-10])
# Central-directory declarations must be checked before inflation/allocation.
large = bytearray((ROOT/'precision.xlsx').read_bytes()); pos = large.index(b'PK\x01\x02')
large[pos+24:pos+28] = (50*1024*1024+1).to_bytes(4,'little')
(ROOT/'expanded-limit.xlsx').write_bytes(large)
print('Generated', len(list(ROOT.iterdir())), 'synthetic fixture files')